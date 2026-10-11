"""Stage 3 - Present: register.json -> live Excel BoQ. All quantities in the BoQ are SUMIFS formulas
over the element register, so changing a classification key, a parameter or a rule cell re-flows."""
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.comments import Comment

D = json.load(open('register.json')); REG, QA = D['register'], D['qa']
U = json.load(open('unified.json')); META = U['meta']
SP = [r for r in U['elements'] if r['ifc_class'] == 'IfcSpace']

F = 'Arial'
f_norm = Font(name=F, size=10); f_b = Font(name=F, size=10, bold=True)
f_h = Font(name=F, size=10, bold=True, color='FFFFFF'); f_title = Font(name=F, size=14, bold=True)
f_in = Font(name=F, size=10, color='0000FF'); f_grp = Font(name=F, size=10, bold=True, color='3B2A6B')
fill_h = PatternFill('solid', fgColor='3B2A6B'); fill_grp = PatternFill('solid', fgColor='E9E4F5')
fill_in = PatternFill('solid', fgColor='FFFF00'); fill_tot = PatternFill('solid', fgColor='F2F2F2')
thin = Side(style='thin', color='BFBFBF'); box = Border(top=thin, bottom=thin, left=thin, right=thin)
NUM = '#,##0.00;-#,##0.00;"-"'; NUM3 = '#,##0.000;-#,##0.000;"-"'; INT = '#,##0;-#,##0;"-"'; PCT = '0.0%;-0.0%;"-"'

wb = Workbook()
def sheet(name, first=False):
    ws = wb.active if first else wb.create_sheet(name); ws.title = name; return ws
def header(ws, row, cols, widths=None):
    for i, c in enumerate(cols, 1):
        x = ws.cell(row=row, column=i, value=c); x.font = f_h; x.fill = fill_h
        x.alignment = Alignment(wrap_text=True, vertical='center'); x.border = box
    if widths:
        for i, w in enumerate(widths, 1): ws.column_dimensions[L(i)].width = w
def style_rows(ws, r0, r1, ncol):
    for r in range(r0, r1 + 1):
        for c in range(1, ncol + 1):
            x = ws.cell(row=r, column=c)
            if x.font == Font() or x.font.name != F: x.font = f_norm
            x.border = box

readme = sheet('ReadMe', first=True); boq = sheet('BoQ'); reg = sheet('Register'); op = sheet('Openings')
mto = sheet('Materials'); rooms = sheet('Rooms'); qa = sheet('QA'); par = sheet('Params')

# ---------------- Params ----------------
par['A1'] = 'Parameters (edit the yellow cells - everything downstream recalculates)'; par['A1'].font = f_title
header(par, 3, ['Parameter', 'Value', 'Unit', 'Used by / source'], [44, 12, 8, 70])
P = [('Steel density', 7850, 'kg/m3', 'Register col O - cross-check of beam mass from volume. Standard value for structural steel.'),
     ('Opening add-back threshold', 0, 'm2', 'Wall areas: openings with W x H <= this value are added back (not deducted). 0 = deduct all openings (model net). Set per your measurement standard.'),
     ('Strip footing depth', 0.3, 'm', 'BoQ 1.02 formwork. From type name "Bearing Footing - 900 x 300".'),
     ('Ground slab thickness', 0.127, 'm', 'BoQ 1.08 edge formwork. From layer set "127mm Slab on Grade".'),
     ('Exterior slab thickness', 0.15, 'm', 'BoQ 8.3 edge formwork. From layer set "150mm Exterior Slab on Grade".')]
for i, (a, b, c, d) in enumerate(P, 4):
    par.cell(row=i, column=1, value=a); x = par.cell(row=i, column=2, value=b); x.font = f_in; x.fill = fill_in
    par.cell(row=i, column=3, value=c); par.cell(row=i, column=4, value=d)
style_rows(par, 4, 3 + len(P), 4)
for i in range(4, 4 + len(P)): par.cell(row=i, column=2).font = f_in
PRM = dict(density='Params!$B$4', thr='Params!$B$5', ftg='Params!$B$6', gs='Params!$B$7', es='Params!$B$8')

# ---------------- Register ----------------
RC = ['GUID', 'IFC class', 'Family', 'Type', 'Storey', 'BoQ key', 'Count', 'Length (m)', 'Area adopted (m2)',
      'Volume adopted (m3)', 'Perimeter (m)', 'Layer-set thickness (mm)', 'Nominal mass (kg/m)', 'Mass (kg)',
      'Mass check: vol x density (kg)', 'Reported area (m2)', 'Reported volume (m3)', 'Geometric volume (m3)',
      'Geom vs reported', 'Quantity source / reasoning', 'Flag', 'Revit Id']
header(reg, 1, RC, [24, 20, 26, 40, 9, 9, 7, 10, 11, 11, 10, 10, 10, 10, 12, 11, 11, 11, 10, 90, 9, 9])
reg.freeze_panes = 'G2'
order = sorted(REG, key=lambda x: ([int(p) if p.isdigit() else 99 for p in x['key'].split('.')] if x['key'][0].isdigit() else [999], x['storey'] or '', x['type']))
for i, x in enumerate(order, 2):
    vals = [x['guid'], x['ifc_class'], x['family'], x['type'], x['storey'], x['key'], x['count'], x['length'], None,
            x['volume'], x['perimeter'], x['t_mm'], x['kg_per_m'], None, None, x['rep_area'], x['rep_vol'], x['geo_vol'],
            None, x['source'], x['flag'], x['revit_id']]
    for c, v in enumerate(vals, 1): reg.cell(row=i, column=c, value=v)
    reg.cell(row=i, column=9, value=f'=IF(L{i}>0,J{i}/(L{i}/1000),"")' if x['area_formula'] else x['area'])
    if x['kg_per_m']:
        reg.cell(row=i, column=14, value=f'=H{i}*M{i}'); reg.cell(row=i, column=15, value=f'=J{i}*{PRM["density"]}')
    if x['rep_vol'] and x['geo_vol']: reg.cell(row=i, column=19, value=f'=(R{i}-Q{i})/Q{i}')
RN = len(order) + 1
style_rows(reg, 2, RN, len(RC))
for r in range(2, RN + 1):
    for c, fmt in ((8, NUM), (9, NUM), (10, NUM3), (11, NUM), (14, INT), (15, INT), (16, NUM), (17, NUM3), (18, NUM3), (19, PCT)):
        reg.cell(row=r, column=c).number_format = fmt
    if reg.cell(row=r, column=21).value in ('Check', 'Minor', 'Resolved'):
        reg.cell(row=r, column=21).fill = PatternFill('solid', fgColor={'Check': 'F4B6B6', 'Minor': 'FCE4C4', 'Resolved': 'D9EAD3'}[reg.cell(row=r, column=21).value])
reg.auto_filter.ref = f'A1:{L(len(RC))}{RN}'
keycol = {x['guid']: x['key'] for x in REG}; typ = {x['guid']: x['type'] for x in REG}

# ---------------- Openings (doors/windows schedule) ----------------
header(op, 1, ['GUID', 'IFC class', 'Type', 'Storey', 'Width (m)', 'Height (m)', 'Area W x H (m2)', 'External', 'Host element type', 'Host BoQ key', 'Own BoQ key'],
       [24, 11, 34, 9, 10, 10, 12, 9, 40, 11, 11])
ops = sorted([x for x in REG if x['ifc_class'] in ('IfcDoor', 'IfcWindow')], key=lambda x: (x['key'], x['storey']))
for i, x in enumerate(ops, 2):
    host = x['host']; hk = keycol.get(host)
    if hk == 'EXCL': hk = '2.3.1'   # skylight hosted by roof's child slab -> roof
    for c, v in enumerate([x['guid'], x['ifc_class'], x['type'], x['storey'], x['width'], x['height'], f'=E{i}*F{i}',
                           'Yes' if x['is_ext'] else 'No', typ.get(host) or 'Live Roof (via roof slab)', hk, x['key']], 1):
        op.cell(row=i, column=c, value=v)
ON = len(ops) + 1
style_rows(op, 2, ON, 11)
for r in range(2, ON + 1):
    for c in (5, 6, 7): op.cell(row=r, column=c).number_format = NUM3 if c < 7 else NUM
op.freeze_panes = 'A2'

# ---------------- Materials (layer take-off + naming dictionary) ----------------
mto['A1'] = 'Material take-off from IFC material layer sets'; mto['A1'].font = f_title
mto['A2'] = 'Layer area = element adopted area (one face per layer); layer volume = area x thickness. A different VIEW of the same elements - not additional to the BoQ.'
mto['A2'].font = Font(name=F, size=9, italic=True)
header(mto, 4, ['BoQ key', 'Element type', 'Category', 'Layer #', 'Material (as in IFC)', 'Normalised material', 'Thickness (mm)', 'Element area (m2)', 'Layer volume (m3)'],
       [9, 44, 9, 7, 44, 36, 11, 13, 13])
seen, rows = set(), []
uel = {r['guid']: r for r in U['elements']}
for x in order:
    m = uel[x['guid']]['material']
    if x['key'] in seen or not m or m['kind'] != 'layerset' or x['count'] == 0: continue
    seen.add(x['key'])
    cat = 'Wall' if x['ifc_class'].startswith('IfcWall') else {'IfcSlab': 'Floor', 'IfcRoof': 'Roof', 'IfcCovering': 'Ceiling'}[x['ifc_class']]
    for j, l in enumerate(m['layers'], 1): rows.append((x['key'], x['type'], cat, j, l['material'], round(l['thickness'] * 1000, 1)))
mats = sorted(set(r[4] for r in rows))
NORM = {m: m for m in mats}; NORM['Concrete'] = 'Concrete - Cast In Situ'
# naming dictionary table (columns K:L)
mto['K4'] = 'Material (as in IFC)'; mto['L4'] = 'Normalised name (edit)'
for c in ('K4', 'L4'): mto[c].font = f_h; mto[c].fill = fill_h; mto[c].border = box
mto.column_dimensions['K'].width = 46; mto.column_dimensions['L'].width = 40
for i, m in enumerate(mats, 5):
    mto.cell(row=i, column=11, value=m).font = f_norm; c = mto.cell(row=i, column=12, value=NORM[m]); c.font = f_in; c.fill = fill_in
    mto.cell(row=i, column=11).border = box; c.border = box
dict_end = 4 + len(mats)
for i, r in enumerate(rows, 5):
    k, t, cat, j, m, th = r
    for c, v in enumerate([k, t, cat, j, m, f'=INDEX($L$5:$L${dict_end},MATCH(E{i},$K$5:$K${dict_end},0))', th,
                           f'=SUMIFS(Register!$I:$I,Register!$F:$F,A{i})', f'=H{i}*G{i}/1000'], 1):
        mto.cell(row=i, column=c, value=v)
ME = 4 + len(rows)
style_rows(mto, 5, ME, 9)
for r in range(5, ME + 1): mto.cell(row=r, column=8).number_format = NUM; mto.cell(row=r, column=9).number_format = NUM3
# summary by normalised material
s0 = ME + 3
mto.cell(row=s0 - 1, column=1, value='Summary by normalised material').font = f_b
for c, h in enumerate(['', 'Normalised material', 'Total layer area (m2)', 'Total volume (m3)'], 1):
    if h: x = mto.cell(row=s0, column=c, value=h); x.font = f_h; x.fill = fill_h; x.border = box
norms = sorted(set(NORM.values()))
for i, n in enumerate(norms, s0 + 1):
    mto.cell(row=i, column=2, value=n)
    mto.cell(row=i, column=3, value=f'=SUMIFS($H$5:$H${ME},$F$5:$F${ME},B{i})').number_format = NUM
    mto.cell(row=i, column=4, value=f'=SUMIFS($I$5:$I${ME},$F$5:$F${ME},B{i})').number_format = NUM3
extra = s0 + 1 + len(norms)
mto.cell(row=extra, column=2, value='Concrete - strip footings (no material in IFC; BoQ 1.01)')
mto.cell(row=extra, column=4, value='=SUMIFS(Register!$J:$J,Register!$F:$F,"1.01")').number_format = NUM3
mto.cell(row=extra + 1, column=2, value='Metal - Steel - 345 MPa, beams (kg - see Register col N)')
mto.cell(row=extra + 1, column=4, value='=SUMIFS(Register!$N:$N,Register!$F:$F,"2.1.*")').number_format = INT
for r in range(s0 + 1, extra + 2):
    for c in (2, 3, 4): mto.cell(row=r, column=c).font = f_norm; mto.cell(row=r, column=c).border = box

# ---------------- Rooms ----------------
header(rooms, 1, ['Number', 'Name', 'Level', 'Category', 'Area (m2)', 'Perimeter (m)', 'Height (m)', 'Volume (m3)', 'GSA BIM area (m2)', 'In NIA? (edit)'],
       [9, 16, 9, 34, 11, 12, 10, 11, 13, 12])
sp = sorted(SP, key=lambda r: (r['storey'], r['number']))
for i, r in enumerate(sp, 2):
    rp = r['reported']
    for c, v in enumerate([r['number'], r['room_name'], r['storey'], r['category'], rp.get('Area'), rp.get('Perimeter'),
                           rp.get('Unbounded Height'), rp.get('Volume'), r['gsa_area'], 'N' if r['number'] == 'R301' else 'Y'], 1):
        rooms.cell(row=i, column=c, value=v)
SN = len(sp) + 1
style_rows(rooms, 2, SN, 10)
for r in range(2, SN + 1):
    for c in (5, 6, 7, 8, 9): rooms.cell(row=r, column=c).number_format = NUM
    rooms.cell(row=r, column=10).font = f_in; rooms.cell(row=r, column=10).fill = fill_in
t0 = SN + 2
rooms.cell(row=t0, column=2, value='Net internal area by level').font = f_b
for i, lv in enumerate(['Level 1', 'Level 2', 'Roof'], t0 + 1):
    rooms.cell(row=i, column=3, value=lv).font = f_norm
    rooms.cell(row=i, column=5, value=f'=SUMIFS($E$2:$E${SN},$C$2:$C${SN},C{i},$J$2:$J${SN},"Y")').number_format = NUM
rooms.cell(row=t0 + 4, column=3, value='Total NIA').font = f_b
rooms.cell(row=t0 + 4, column=5, value=f'=SUM(E{t0+1}:E{t0+3})').number_format = NUM
rooms.cell(row=t0 + 4, column=5).font = f_b
rooms.cell(row=t0 + 5, column=2, value='R301 is the roof surface space (green roof), excluded by default. Stair spaces A105/B105 are double height and counted once on Level 1.').font = Font(name=F, size=9, italic=True)
rooms.freeze_panes = 'A2'

# ---------------- QA ----------------
qa['A1'] = 'QA log - what the checks found and what was done about it'; qa['A1'].font = f_title
header(qa, 3, ['Status', 'Element(s)', 'Finding', 'Action / decision'], [13, 34, 80, 80])
for i, row in enumerate(QA, 4):
    for c, v in enumerate(row, 1):
        x = qa.cell(row=i, column=c, value=v); x.alignment = Alignment(wrap_text=True, vertical='top')
style_rows(qa, 4, 3 + len(QA), 4)
qa.cell(row=4 + len(QA), column=1, value='Data quality').font = f_norm
qa.cell(row=4 + len(QA), column=2, value='Material names').font = f_norm
qa.cell(row=4 + len(QA), column=3, value='Two names for the same material: "Concrete" (127 mm slab) and "Concrete - Cast In Situ".').font = f_norm
qa.cell(row=4 + len(QA), column=4, value='Merged via the editable naming dictionary on the Materials sheet (cols K:L).').font = f_norm
for c in range(1, 5): qa.cell(row=4 + len(QA), column=c).border = box; qa.cell(row=4 + len(QA), column=c).alignment = Alignment(wrap_text=True, vertical='top')

# ---------------- BoQ ----------------
def S(col, key): return f'SUMIFS(Register!${col}:${col},Register!$F:$F,"{key}")'
def WALL(key): return f'={S("I", key)}+SUMIFS(Openings!$G:$G,Openings!$J:$J,"{key}",Openings!$G:$G,"<="&{PRM["thr"]})'
def NR(key): return f'={S("G", key)}'
src_rev = 'Reported (Revit) volume, geometry-verified'
src_wall = 'Net area = volume / layer-set thickness (all openings deducted), + add-back of openings <= threshold (Params)'
B = [
 ('G', '1', 'SUBSTRUCTURE'),
 ('I', '1.01', 'In-situ concrete strip footings, 900 x 300 mm', 'm3', f'={S("J","1.01")}', src_rev + '; material assumed concrete (none in IFC)'),
 ('I', '1.02', 'Formwork to vertical sides of strip footings (derived: 2 x length x depth)', 'm2', f'=2*{S("H","1.01")}*{PRM["ftg"]}', 'Derived from reported footing length'),
 ('I', '1.03', 'In-situ concrete foundation walls, 417 mm thick', 'm3', f'={S("J","1.03")}', src_rev),
 ('I', '1.04', 'In-situ concrete foundation walls, 435 mm thick', 'm3', f'={S("J","1.04")}', src_rev),
 ('I', '1.05', 'Formwork to both faces of foundation walls (derived: 2 x net area)', 'm2', f'=2*({S("I","1.03")}+{S("I","1.04")})', 'Derived from adopted wall area'),
 ('I', '1.06', 'In-situ concrete ground-bearing slab, 127 mm thick', 'm3', f'={S("J","1.06")}', src_rev),
 ('I', '1.07', 'Ground-bearing slab 127 mm - surface area (for DPM, mesh, power-floating)', 'm2', f'={S("I","1.06")}', 'Reported area'),
 ('I', '1.08', 'Edge formwork to ground slab, 127 mm high (derived: perimeter x thickness)', 'm2', f'={S("K","1.06")}*{PRM["gs"]}', 'Derived from reported perimeter'),
 ('G', '2.1', 'FRAME'),
 ('I', '2.1.1', 'Structural steel beam W310x60, grade 345 MPa', 't', f'={S("N","2.1.1")}/1000', 'Reported length x nominal 60 kg/m (volume x density in Register as check)'),
 ('I', '2.1.2', 'Structural steel beam W410x60, grade 345 MPa', 't', f'={S("N","2.1.2")}/1000', 'Reported length x nominal 60 kg/m'),
 ('G', '2.2', 'UPPER FLOORS'),
 ('I', '2.2.1', 'Timber joist floor: 286 mm dimensional lumber joists, 19 mm plywood subfloor (305 mm overall)', 'm2', f'={S("I","2.2.1")}', 'Reported area, V/t-consistent'),
 ('G', '2.3', 'ROOF'),
 ('I', '2.3.1', 'Flat live (green) roof: 64 mm growing medium, root barrier, EPDM membrane, 76 mm rigid insulation, 19 mm plywood, 286 mm timber joists (457 mm)', 'm2', f'={S("I","2.3.1")}', 'Reported area on IfcRoof; geometry from child slab'),
 ('I', '2.3.2', 'Roof light, 1180 x 1170 mm', 'nr', NR('2.3.2'), 'Count'),
 ('G', '2.4', 'STAIRS AND RAMPS'),
 ('I', '2.4.1', 'Straight-flight stair, 16 risers @ 193.75 mm, 250 mm going, 1000 mm wide, Level 1 to Level 2 (material not defined in model)', 'nr', NR('2.4.1'), 'Count; geometry and riser data from Pset_StairFlightCommon'),
 ('I', '2.4.2', 'Handrail, 900 mm high', 'm', f'={S("H","2.4.2")}', 'Reported length'),
 ('I', '2.4.3', 'Guard rail, 1100 mm high', 'm', f'={S("H","2.4.3")}', 'Reported length'),
 ('G', '2.5', 'EXTERNAL WALLS'),
 ('I', '2.5.1', 'Cavity wall: 92 mm brick, 25 mm cavity, 50 mm rigid insulation, 193 mm concrete block, 41 mm metal stud, 16 mm plasterboard (417 mm)', 'm2', WALL('2.5.1'), src_wall),
 ('I', '2.5.2', 'Ditto, parapet above roof level', 'm2', WALL('2.5.2'), src_wall),
 ('G', '2.6', 'WINDOWS AND EXTERNAL DOORS'),
 ('I', '2.6.1', 'Fixed window, 4835 x 2420 mm', 'nr', NR('2.6.1'), 'Count'),
 ('I', '2.6.2', 'Fixed window, 2800 x 2410 mm', 'nr', NR('2.6.2'), 'Count'),
 ('I', '2.6.3', 'Fixed window, 750 x 2200 mm', 'nr', NR('2.6.3'), 'Count'),
 ('I', '2.6.4', 'Fixed window, 819 x 759 mm', 'nr', NR('2.6.4'), 'Count'),
 ('I', '2.6.5', 'Casement window, 819 x 759 mm', 'nr', NR('2.6.5'), 'Count'),
 ('I', '2.6.6', 'External door, single flush, 1250 x 2010 mm', 'nr', NR('2.6.6'), 'Count'),
 ('I', '2.6.7', 'External door, single glazed, 813 x 2420 mm', 'nr', NR('2.6.7'), 'Count'),
 ('G', '2.7', 'INTERNAL WALLS AND PARTITIONS'),
 ('I', '2.7.1', 'Party wall: 2 x 193 mm concrete block, 50 mm cavity, 41 mm metal stud + 16 mm plasterboard each side (550 mm)', 'm2', WALL('2.7.1'), src_wall + '; see QA (junction overlap, surface-model repair)'),
 ('I', '2.7.2', 'Partition: 92 mm metal stud, 16 mm plasterboard both sides (124 mm)', 'm2', WALL('2.7.2'), src_wall),
 ('I', '2.7.3', 'Plumbing wall: 152 mm metal stud, 16 mm plasterboard both sides (184 mm)', 'm2', WALL('2.7.3'), src_wall),
 ('I', '2.7.4', 'Wall lining: 136 mm metal stud furring, 16 mm plasterboard one side (152 mm)', 'm2', WALL('2.7.4'), src_wall),
 ('I', '2.7.5', 'Wall lining: 38 mm metal stud furring, 16 mm plasterboard one side (54 mm)', 'm2', WALL('2.7.5'), src_wall),
 ('G', '2.8', 'INTERNAL DOORS'),
 ('I', '2.8.1', 'Internal door, single flush, 762 x 2032 mm', 'nr', NR('2.8.1'), 'Count'),
 ('I', '2.8.2', 'Internal door, single flush, 864 x 2032 mm', 'nr', NR('2.8.2'), 'Count'),
 ('G', '3.1', 'WALL FINISHES'),
 ('I', '3.1.1', 'Emulsion paint to plasterboard wall surfaces (derived: sum of plasterboard layer faces on walls)', 'm2', '=SUMIFS(Materials!$H:$H,Materials!$E:$E,"Plasterboard",Materials!$C:$C,"Wall")', 'Derived from material layers'),
 ('G', '3.2', 'FLOOR FINISHES'),
 ('I', '3.2.1', 'Timber floor finish: 16 mm wood flooring on 3 mm underlay', 'm2', f'={S("I","3.2.1")}', 'Reported area, V/t-consistent'),
 ('I', '3.2.2', 'Ceramic floor tiling: 10 mm tile on 3 mm bed/grout', 'm2', f'={S("I","3.2.2")}', 'Reported area, V/t-consistent'),
 ('G', '3.3', 'CEILING FINISHES'),
 ('I', '3.3.1', 'Suspended ceiling: 12 mm plasterboard on 45 mm metal studs', 'm2', f'={S("I","3.3.1")}', 'Reported area; ~3.2 m2 possible double-modelling (QA)'),
 ('I', '3.3.2', 'Emulsion paint to plasterboard ceilings (derived)', 'm2', '=__ROW_3.3.1__', 'Derived = 3.3.1'),
 ('G', '4', 'FITTINGS, FURNISHINGS AND EQUIPMENT'),
 ('I', '4.1.1', 'Kitchen base cabinet, double door + 2 drawers, 1000 mm', 'nr', NR('4.1.1'), 'Count'),
 ('I', '4.1.2', 'Kitchen wall cabinet, double door, 1000 mm', 'nr', NR('4.1.2'), 'Count (two Revit type names merged)'),
 ('I', '4.1.3', 'Tall cabinet, single door, 800 mm', 'nr', NR('4.1.3'), 'Count'),
 ('I', '4.1.4', 'Worktop with sink cut-out, 600 mm deep', 'nr', NR('4.1.4'), 'Count (length not exported)'),
 ('I', '4.1.5', 'Worktop, 600 mm deep', 'nr', NR('4.1.5'), 'Count (length not exported)'),
 ('I', '4.1.6', 'Vanity unit with basin, 450 x 450 mm', 'nr', NR('4.1.6'), 'Count'),
 ('I', '4.1.7', 'Vanity unit with basin, 650 x 450 mm', 'nr', NR('4.1.7'), 'Count'),
 ('I', '4.2.1', 'Loose furniture (FF&E, for information): queen bed 1525 x 2007', 'nr', NR('4.2.1'), 'Count - normally client-supplied'),
 ('I', '4.2.2', 'Loose furniture (FF&E): king bed 1981 x 2032', 'nr', NR('4.2.2'), 'Count'),
 ('I', '4.2.3', 'Loose furniture (FF&E): sofa 1830', 'nr', NR('4.2.3'), 'Count'),
 ('I', '4.2.4', 'Loose furniture (FF&E): coffee table 610 x 610', 'nr', NR('4.2.4'), 'Count'),
 ('I', '4.2.5', 'Loose furniture (FF&E): coffee table 915 x 1830', 'nr', NR('4.2.5'), 'Count'),
 ('G', '8', 'EXTERNAL WORKS'),
 ('I', '8.1', 'In-situ concrete external slab on grade, 150 mm thick', 'm3', f'={S("J","8.1")}', src_rev),
 ('I', '8.2', 'External slab 150 mm - surface area', 'm2', f'={S("I","8.1")}', 'Reported area'),
 ('I', '8.3', 'Edge formwork to external slab, 150 mm high (derived)', 'm2', f'={S("K","8.1")}*{PRM["es"]}', 'Derived from reported perimeter'),
]
boq['A1'] = f"Bill of Quantities - {META['source']}"; boq['A1'].font = f_title
boq['A2'] = 'Elemental grouping modelled on NRM1. Quantities are live formulas over the Register. Enter rates in the yellow column.'
boq['A2'].font = Font(name=F, size=9, italic=True)
header(boq, 4, ['Ref', 'Description', 'Unit', 'Quantity', 'Rate', 'Amount', 'Measurement basis / source'], [8, 92, 6, 12, 11, 14, 70])
boq.freeze_panes = 'A5'
r = 5; grp_rows = []; cur = None; rowref = {}
def close_group():
    global r
    if cur is None: return
    g0, gname = cur
    boq.cell(row=r, column=2, value=f'Subtotal {gname.title()}').font = f_b
    c = boq.cell(row=r, column=6, value=f'=SUM(F{g0+1}:F{r-1})'); c.font = f_b; c.number_format = NUM
    for cc in range(1, 8): boq.cell(row=r, column=cc).fill = fill_tot; boq.cell(row=r, column=cc).border = box
    grp_rows.append(r); r += 2
for item in B:
    if item[0] == 'G':
        close_group(); cur = (r, item[2])
        boq.cell(row=r, column=1, value=item[1]).font = f_grp; boq.cell(row=r, column=2, value=item[2]).font = f_grp
        for cc in range(1, 8): boq.cell(row=r, column=cc).fill = fill_grp; boq.cell(row=r, column=cc).border = box
        r += 1; continue
    _, ref, desc, unit, qf, basis = item
    rowref[ref] = r
    qf = qf.replace('__ROW_3.3.1__', f'D{rowref.get("3.3.1", 0)}')
    vals = [ref, desc, unit, qf, None, f'=D{r}*E{r}', basis]
    for c, v in enumerate(vals, 1):
        x = boq.cell(row=r, column=c, value=v); x.font = f_norm; x.border = box
        x.alignment = Alignment(wrap_text=c in (2, 7), vertical='top')
    boq.cell(row=r, column=4).number_format = INT if unit == 'nr' else (NUM3 if unit == 't' else NUM)
    e = boq.cell(row=r, column=5); e.fill = fill_in; e.font = f_in; e.number_format = NUM
    boq.cell(row=r, column=6).number_format = NUM
    r += 1
close_group()
boq.cell(row=r, column=2, value='TOTAL').font = Font(name=F, size=11, bold=True)
c = boq.cell(row=r, column=6, value='=' + '+'.join(f'F{x}' for x in grp_rows)); c.font = Font(name=F, size=11, bold=True); c.number_format = NUM
for cc in range(1, 8): boq.cell(row=r, column=cc).border = Border(top=Side(style='double'))
boq.cell(row=r + 2, column=2, value='Excluded / not in model: MEP, reinforcement, skirtings, sills, ironmongery, sitework beyond the two external slabs.').font = Font(name=F, size=9, italic=True)

# ---------------- ReadMe ----------------
readme.column_dimensions['A'].width = 3; readme.column_dimensions['B'].width = 26; readme.column_dimensions['C'].width = 110
readme['B1'] = 'IFC -> Bill of Quantities'; readme['B1'].font = f_title
info = [('Source file', META['source']), ('Schema', META['schema']), ('Authoring tool', META['authoring_app']),
        ('Units', 'Metres / m2 / m3 (from IfcUnitAssignment)'), ('Storeys', ', '.join(f"{k} ({v:+.2f})" for k, v in sorted(META['storeys'].items(), key=lambda kv: kv[1]))),
        ('Elements measured', f"{sum(1 for x in REG if x['count'])} physical elements + {len(SP)} spaces"), ('', ''),
        ('Pipeline', '1. Extract (extract.py): IFC -> unified.json. Per element: reported quantities, geometry-measured volume/area, type, storey, material layers, host/aggregate links.'),
        ('', '2. Verify (extract.py + overlap.py): every solid re-meshed and measured; reported vs geometric volume compared at 1% tolerance; pairwise solid intersections to catch double counting.'),
        ('', '3. Classify & measure (classify.py): rule table maps class + type (+ storey) to a BoQ item key; adopts a quantity per element with a stated source.'),
        ('', '4. Present (build_boq.py): this workbook. BoQ lines are SUMIFS over the Register, so keys, parameters and the naming dictionary can be edited live.'),
        ('', ''),
        ('Sheets', 'BoQ - priced-bill layout, rates column blank.  Register - one row per element, with the reasoning behind each number.  Openings - door/window schedule with host walls.'),
        ('', 'Materials - layer-by-layer take-off + editable naming dictionary.  Rooms - space schedule and NIA.  QA - findings and decisions.  Params - editable rule parameters.'),
        ('', ''),
        ('Colour legend', 'Yellow fill + blue text = input you can edit (rates, parameters, NIA flags, material names). Black = formula or extracted value.'),
        ('Register flags', 'Resolved (green) = discrepancy found and explained. Minor (orange) = small known deviation, handled by a different measure. Check (red) = needs a human.')]
for i, (a, b) in enumerate(info, 3):
    readme.cell(row=i, column=2, value=a).font = f_b; x = readme.cell(row=i, column=3, value=b); x.font = f_norm; x.alignment = Alignment(wrap_text=True)

for ws in wb.worksheets: ws.sheet_view.showGridLines = ws.title in ('Register', 'Openings')
wb.save('/mnt/user-data/outputs/Duplex_BoQ.xlsx'); print('saved; BoQ rows to', r)
