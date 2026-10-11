"""
export_xlsx.py - BoqResult -> Excel workbook.

Sheets
  BoQ          priced-bill layout grouped by the first classification system; Rate is an
               input column, Amount and subtotals are live formulas
  Calculation  every line with the formula, parameters and each element's contribution
  Lines        flat table of all lines with codes in every system (for filtering / pivots)
  Log          the calculation log
  Questions    open questions, warnings and the parser's verification checks
  Inputs       model, rule set (with parameter values), dictionaries, classifications, decisions
"""
from __future__ import annotations

import io
import json
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

F = "Arial"
H_FONT = Font(name=F, size=10, bold=True, color="FFFFFF")
H_FILL = PatternFill("solid", fgColor="3B2A6B")
SEC_FILL = PatternFill("solid", fgColor="E9E4F5")
SUB_FILL = PatternFill("solid", fgColor="F2F2F2")
FLAG_FILL = PatternFill("solid", fgColor="FFF2CC")
IN_FILL = PatternFill("solid", fgColor="FFFF00")
NORMAL = Font(name=F, size=10)
BOLD = Font(name=F, size=10, bold=True)
INPUT = Font(name=F, size=10, color="0000FF")
TITLE = Font(name=F, size=14, bold=True)
SMALL = Font(name=F, size=9, italic=True, color="595959")
THIN = Side(style="thin", color="D0D0D0")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
NUM = '#,##0.00;-#,##0.00;"-"'
NUM3 = '#,##0.000;-#,##0.000;"-"'
WRAP = Alignment(wrap_text=True, vertical="top")


def _header(ws, row, cols, widths):
    for i, (c, w) in enumerate(zip(cols, widths), 1):
        x = ws.cell(row=row, column=i, value=c)
        x.font, x.fill, x.border = H_FONT, H_FILL, BOX
        x.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w


def _row(ws, row, values, font=NORMAL, fill=None, wrap_cols=()):
    for i, v in enumerate(values, 1):
        x = ws.cell(row=row, column=i, value=v)
        x.font, x.border = font, BOX
        if fill:
            x.fill = fill
        if i in wrap_cols:
            x.alignment = WRAP


def _qfmt(unit):
    return NUM3 if unit in ("t",) else ("0" if unit in ("nr", "item") else NUM)


def export(result: dict, meta: dict) -> bytes:
    wb = Workbook()
    systems = meta.get("classifications", [])          # [{id, name}]
    primary = systems[0]["id"] if systems else None
    lines = result["lines"]

    # ------------------------------------------------------------------ BoQ
    ws = wb.active
    ws.title = "BoQ"
    ws["A1"] = f"Bill of Quantities - {meta.get('modelName', '')}"
    ws["A1"].font = TITLE
    ws["A2"] = (f"Rules: {result['ruleSet']['name']} ({result['ruleSet']['meta'].get('chain', result['ruleSet']['id'])})  |  "
                f"Grouping: {systems[0]['name'] if systems else 'none'}  |  Status: {result['status']}  |  "
                f"Generated {meta.get('generatedAt', datetime.now().isoformat(timespec='seconds'))}")
    ws["A2"].font = SMALL
    ws["A3"] = "Enter rates in the yellow column. Amounts and totals are formulas. Yellow rows are flagged: quantity to be confirmed."
    ws["A3"].font = SMALL
    _header(ws, 5, ["Ref", "Description", "Unit", "Quantity", "Rate", "Amount", "Source", "Notes"],
            [10, 70, 7, 12, 11, 14, 10, 50])
    ws.freeze_panes = "A6"
    groups: dict[tuple, list] = {}
    for l in lines:
        c = l["codes"].get(primary) if primary else None
        top = (c["path"][0]["code"], c["path"][0]["title"]) if c else ("~", "Unclassified")
        groups.setdefault(top, []).append(l)
    r = 6
    subtotal_cells = []
    for (code, title), ls in sorted(groups.items(), key=lambda kv: kv[0][0]):
        _row(ws, r, [code if code != "~" else "", title, "", "", "", "", "", ""], BOLD, SEC_FILL)
        r += 1
        first = r
        for l in ls:
            c = l["codes"].get(primary) if primary else None
            flagged = l["quantity"] is None
            _row(ws, r, [c["code"] if c else "", l["description"], l["unit"], l["quantity"], None,
                         f"=IF(OR(D{r}=\"\",E{r}=\"\"),0,D{r}*E{r})", l["source"], "; ".join(l.get("notes", []))],
                 NORMAL, FLAG_FILL if flagged else None, wrap_cols=(2, 8))
            ws.cell(row=r, column=4).number_format = _qfmt(l["unit"])
            ws.cell(row=r, column=5).fill = IN_FILL
            ws.cell(row=r, column=5).font = INPUT
            ws.cell(row=r, column=5).number_format = NUM
            ws.cell(row=r, column=6).number_format = NUM
            r += 1
        _row(ws, r, ["", f"Subtotal {title}", "", "", "", f"=SUM(F{first}:F{r - 1})", "", ""], BOLD, SUB_FILL)
        ws.cell(row=r, column=6).number_format = NUM
        subtotal_cells.append(f"F{r}")
        r += 2
    _row(ws, r, ["", "TOTAL", "", "", "", "=" + ("+".join(subtotal_cells) if subtotal_cells else "0"), "", ""], Font(name=F, size=11, bold=True))
    ws.cell(row=r, column=6).number_format = NUM

    # ------------------------------------------------------------------ Calculation
    wc = wb.create_sheet("Calculation")
    wc["A1"] = "How every quantity was calculated"
    wc["A1"].font = TITLE
    wc["A2"] = "Each line: its formula and parameters, then every element's contribution with the values used."
    wc["A2"].font = SMALL
    _header(wc, 4, ["Line / element", "Quantity", "Unit", "Formula / explanation", "Element GUID"], [60, 12, 7, 110, 26])
    wc.freeze_panes = "A5"
    r = 5
    for l in lines:
        params = ", ".join(f"{k}={json.dumps(v)}" for k, v in l["trace"].get("params", {}).items())
        _row(wc, r, [l["description"], l["quantity"], l["unit"],
                     f"{l['trace']['formula']}" + (f"   [{params}]" if params else ""), ""], BOLD, SEC_FILL, wrap_cols=(1, 4))
        wc.cell(row=r, column=2).number_format = _qfmt(l["unit"])
        r += 1
        for e in l["trace"]["elements"]:
            _row(wc, r, ["    " + (e.get("name") or e["guid"]), e["quantity"], l["unit"], e.get("note"), e["guid"]],
                 NORMAL, None, wrap_cols=(1, 4))
            wc.cell(row=r, column=2).number_format = _qfmt(l["unit"]) if l["unit"] not in ("nr", "item") else "0.###"
            r += 1
        if l.get("trace", {}).get("sourceLine"):
            _row(wc, r, ["    derived from line: " + l["trace"]["sourceLine"], "", "", "", ""], SMALL)
            r += 1

    # ------------------------------------------------------------------ Lines
    wl = wb.create_sheet("Lines")
    cols = ["Key", "Rule", "Kind", "Description", "Unit", "Quantity", "Source"]
    widths = [40, 18, 9, 60, 7, 12, 10]
    for s in systems:
        cols += [f"{s['name']} code", f"{s['name']} title"]
        widths += [12, 30]
    cols += ["Formula", "Elements", "Storey", "Notes"]
    widths += [50, 9, 16, 40]
    _header(wl, 1, cols, widths)
    wl.freeze_panes = "A2"
    for i, l in enumerate(lines, 2):
        vals = [l["key"], l["ruleId"], l["ruleKind"], l["description"], l["unit"], l["quantity"], l["source"]]
        for s in systems:
            c = l["codes"].get(s["id"])
            vals += [c["code"] if c else "", c["title"] if c else ""]
        loc = l.get("location") or {}
        vals += [l["trace"]["formula"], len(l["trace"]["elements"]), loc.get("storey") or ", ".join(loc.get("storeys", [])),
                 "; ".join(l.get("notes", []))]
        _row(wl, i, vals)
        wl.cell(row=i, column=6).number_format = _qfmt(l["unit"])
    wl.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(lines) + 1}"

    # ------------------------------------------------------------------ Log
    wg = wb.create_sheet("Log")
    _header(wg, 1, ["#", "Stage", "Level", "Message", "Rule", "Line", "Element"], [6, 10, 9, 120, 18, 30, 24])
    wg.freeze_panes = "A2"
    for i, e in enumerate(result["log"], 2):
        _row(wg, i, [e["seq"], e["stage"], e["level"], e["message"], e.get("ruleId"), e.get("lineKey"), e.get("guid")],
             NORMAL, FLAG_FILL if e["level"] == "warning" else None)
    wg.auto_filter.ref = f"A1:G{len(result['log']) + 1}"

    # ------------------------------------------------------------------ Questions
    wq = wb.create_sheet("Questions")
    _header(wq, 1, ["Type", "Id / code", "Blocking", "Message", "Elements"], [16, 40, 9, 110, 10])
    r = 2
    for q in result["questions"]:
        _row(wq, r, ["question: " + q["kind"], q["id"], "yes" if q["blocking"] else "no", q["message"], len(q.get("guids") or [])],
             NORMAL, FLAG_FILL if q["blocking"] else None, wrap_cols=(4,))
        r += 1
    for w in result["warnings"]:
        _row(wq, r, ["warning", w["code"], "", w["message"], len(w.get("guids") or [])], NORMAL, None, wrap_cols=(4,))
        r += 1
    for c in result["checks"]:
        _row(wq, r, ["check: " + c["severity"], c["code"], "", c["message"], len(c.get("elementIds") or [])], NORMAL, None, wrap_cols=(4,))
        r += 1

    # ------------------------------------------------------------------ Inputs
    wi = wb.create_sheet("Inputs")
    wi["A1"] = "Inputs of this run"
    wi["A1"].font = TITLE
    rs = result["ruleSet"]
    info = [("Model", meta.get("modelName")), ("Model id (sha256)", meta.get("modelId")),
            ("Rule set", f"{rs['name']} ({rs['id']} v{rs['version']})"), ("Rule set chain", rs["meta"].get("chain")),
            ("Dictionaries", ", ".join(meta.get("dictionaries", []))),
            ("Classifications", ", ".join(s["name"] for s in systems)),
            ("Decisions applied", result["stats"]["decisionsApplied"]),
            ("Elements", result["stats"]["elements"]), ("Measured", result["stats"]["measuredElements"]),
            ("Excluded", result["stats"]["excludedElements"]), ("Unmatched", result["stats"]["unmatchedElements"]),
            ("Engine time (ms)", result["stats"]["durationMs"])]
    for i, (k, v) in enumerate(info, 3):
        wi.cell(row=i, column=1, value=k).font = BOLD
        wi.cell(row=i, column=2, value=v).font = NORMAL
    r = 4 + len(info)
    _header(wi, r, ["Parameter", "Value", "Unit", "Label"], [28, 40, 8, 70])
    for name, p in rs["params"].items():
        r += 1
        _row(wi, r, [name, json.dumps(p.get("value")) if not isinstance(p.get("value"), (int, float, str)) else p.get("value"),
                     p.get("unit", ""), p.get("label", "")])
    if meta.get("decisions"):
        r += 2
        _header(wi, r, ["Decision (question id)", "Answer", "", ""], [28, 40, 8, 70])
        for d in meta["decisions"]:
            r += 1
            _row(wi, r, [d["questionId"], json.dumps(d.get("answer")), "", d.get("note", "")])

    for sheet in wb.worksheets:
        sheet.sheet_view.showGridLines = False
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
